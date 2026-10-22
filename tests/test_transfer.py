import csv
import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from datakyt.services import tabular
from datakyt.services.errors import InvalidInputError


def upload(client, entity, content, filename="data.csv", dry_run=True):
    return client.post(
        f"/import/{entity}",
        params={"dry_run": str(dry_run).lower()},
        files={"file": (filename, content)},
    )


def test_export_csv(client, factory):
    factory.office("Kyiv")
    factory.office("Lviv")
    response = client.get("/export/offices.csv")
    assert response.status_code == 200
    assert 'filename="offices.csv"' in response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert [r["city"] for r in rows] == ["Kyiv", "Lviv"]
    assert set(rows[0]) == {"id", "city", "created_at", "updated_at"}


def test_export_xlsx_skips_deleted_rows(client, factory):
    kept = factory.equipment()
    gone = factory.equipment()
    client.delete(f"/equipment/{gone.id}")
    response = client.get("/export/equipment.xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    header = [c.value for c in sheet[1]]
    assert "deleted_at" not in header
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][header.index("inventory_number")] == kept.inventory_number
    assert rows[0][header.index("cost")] == 1000.0


def test_export_unknown_entity_or_format(client):
    assert client.get("/export/planets.csv").status_code == 404
    assert client.get("/export/offices.json").status_code == 422


def test_import_dry_run_reports_row_errors(client, factory):
    factory.office("Kyiv")
    content = b"city\nOdesa\nKyiv\n\nLviv\n,\n"
    response = upload(client, "offices", content)
    assert response.status_code == 200
    result = response.json()
    assert result["dry_run"] is True
    assert result["imported"] == 0
    assert [e["row"] for e in result["errors"]] == [2, 4]
    assert "uq_office_city" in result["errors"][0]["errors"][0]
    assert "city" in result["errors"][1]["errors"][0]
    assert client.get("/offices").json()["total"] == 1


def test_import_commits_only_when_clean(client, factory):
    project = factory.project()
    office = factory.office()
    header = "full_name,email,phone,project_id,office_id\n"
    good = f"Anna,anna@example.com,,{project.id},{office.id}\nIvan,ivan@example.com,+380501112233,{project.id},{office.id}\n"
    dry = upload(client, "employees", (header + good).encode()).json()
    assert dry["errors"] == [] and dry["imported"] == 0
    assert client.get("/employees").json()["total"] == 0

    done = upload(client, "employees", (header + good).encode(), dry_run=False).json()
    assert done["imported"] == 2
    assert client.get("/employees").json()["total"] == 2

    bad = f"Olena,olena@example.com,,{project.id},999999\n"
    failed = upload(client, "employees", (header + bad).encode(), dry_run=False).json()
    assert failed["imported"] == 0
    assert failed["errors"][0]["row"] == 1
    assert client.get("/employees").json()["total"] == 2


def test_import_ignores_exported_metadata_columns(client, factory):
    factory.office("Kyiv")
    exported = client.get("/export/offices.csv").content
    client.delete(f"/offices/{client.get('/offices').json()['items'][0]['id']}")
    result = upload(client, "offices", exported, dry_run=False).json()
    assert result["errors"] == []
    assert result["imported"] == 1


def test_import_xlsx(client, factory):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "inventory_number",
            "name",
            "serial_number",
            "warranty_months",
            "cost",
            "purchase_date",
            "equipment_type_id",
            "office_id",
            None,
        ]
    )
    type_id, office_id = factory.equipment_type().id, factory.office().id
    sheet.append(["EQ-9", "Monitor", "SN9", 24, 199.99, date(2026, 1, 2), type_id, office_id, "x"])
    sheet.append([None] * 9)
    out = io.BytesIO()
    workbook.save(out)
    result = upload(client, "equipment", out.getvalue(), "items.xlsx", dry_run=False).json()
    assert result == {"dry_run": False, "total_rows": 1, "imported": 1, "errors": []}
    item = client.get("/equipment").json()["items"][0]
    assert item["purchase_date"] == "2026-01-02"
    assert item["cost"] == "199.99"


def test_import_rejects_bad_files(client):
    assert upload(client, "offices", b"city\n", "data.txt").status_code == 422
    assert upload(client, "offices", b"not a zip", "data.xlsx").status_code == 422
    assert upload(client, "offices", "city\nKyiv\n".encode("utf-16"), "data.csv").status_code == 422
    assert upload(client, "planets", b"city\n").status_code == 404


def test_import_limits(client, monkeypatch):
    from datakyt.api.routers import transfer as transfer_router
    from datakyt.services import transfer

    monkeypatch.setattr(transfer, "MAX_IMPORT_ROWS", 1)
    assert upload(client, "offices", b"city\nA\nB\n").status_code == 422
    monkeypatch.setattr(transfer_router, "MAX_UPLOAD_BYTES", 5)
    assert upload(client, "offices", b"city\nA\n").status_code == 422


def test_csv_neutralizes_formulas():
    out = tabular.write_table([{"a": "=HYPERLINK(1)", "b": -5}], ["a", "b"], "csv").decode()
    assert out.splitlines()[1] == "'=HYPERLINK(1),-5"


def test_read_empty_xlsx():
    workbook = Workbook()
    out = io.BytesIO()
    workbook.save(out)
    assert tabular.read_table(out.getvalue(), "xlsx") == []
    with pytest.raises(InvalidInputError):
        tabular.read_table(b"\xff\xfe", "csv")
