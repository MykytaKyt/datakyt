import csv
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from datakyt.services import assignments, reports


def test_inventory_report(client, factory):
    office = factory.office("Kyiv")
    laptop = factory.equipment_type("Laptop")
    for _ in range(2):
        factory.equipment(office_id=office.id, equipment_type_id=laptop.id)
    factory.equipment(office_id=office.id, equipment_type_id=laptop.id, status="broken")
    gone = factory.equipment(office_id=office.id, equipment_type_id=laptop.id)
    client.delete(f"/equipment/{gone.id}")

    rows = client.get("/reports/inventory").json()
    # Statuses sort in enum declaration order.
    assert rows == [
        {"office": "Kyiv", "type": "Laptop", "status": "reserved", "count": 2},
        {"office": "Kyiv", "type": "Laptop", "status": "broken", "count": 1},
    ]


def test_employee_report(client, factory, session):
    employee = factory.employee()
    item = factory.equipment(name="ThinkPad")
    desk = factory.furniture(name="Desk")
    lic = factory.license()
    returned = factory.equipment()
    assignments.issue_equipment(session, item.id, employee.id, None, date(2026, 2, 1))
    assignments.issue_furniture(session, desk.id, employee.id, None, date(2026, 2, 2))
    assignments.assign_license(session, lic.id, employee.id, None, date(2026, 2, 3))
    assignments.issue_equipment(session, returned.id, employee.id, None, date(2026, 2, 1))
    assignments.return_equipment(session, returned.id, None, date(2026, 2, 5))

    rows = client.get(f"/reports/employee/{employee.id}").json()
    assert [(r["category"], r["item_id"]) for r in rows] == [
        ("equipment", item.id),
        ("furniture", desk.id),
        ("license", lic.id),
    ]
    assert rows[0]["name"] == "ThinkPad"
    assert rows[0]["issued_on"] == "2026-02-01"
    assert client.get("/reports/employee/999999").status_code == 404


def test_expiring_licenses(session, factory):
    today = date(2026, 10, 1)
    soon = factory.license(purchased_on=date(2025, 10, 1), expires_on=today + timedelta(days=10))
    factory.license(purchased_on=date(2025, 10, 1), expires_on=today + timedelta(days=90))
    factory.license(purchased_on=date(2025, 1, 1), expires_on=today - timedelta(days=1))
    factory.license(expires_on=None)
    assignments.assign_license(session, soon.id, factory.employee().id, None, date(2025, 10, 2))

    report = reports.expiring_licenses(session, 30, today=today)
    assert [r["license_id"] for r in report.rows] == [soon.id]
    assert report.rows[0]["days_left"] == 10
    assert report.rows[0]["seats_used"] == 1
    assert report.columns[-1] == "days_left"


def test_expiring_warranties(session, factory):
    today = date(2026, 10, 1)
    equipment = factory.equipment(purchase_date=date(2025, 10, 15), warranty_months=12)
    factory.equipment(purchase_date=date(2025, 10, 15), warranty_months=24)
    factory.equipment(purchase_date=date(2024, 1, 1), warranty_months=12)
    desk = factory.furniture(purchase_date=date(2021, 10, 20), warranty_months=60)

    report = reports.expiring_warranties(session, 30, today=today)
    assert [(r["kind"], r["item_id"]) for r in report.rows] == [
        ("equipment", equipment.id),
        ("furniture", desk.id),
    ]
    assert report.rows[0]["warranty_until"] == date(2026, 10, 15)
    assert report.rows[0]["days_left"] == 14


def test_report_endpoints_with_days(client):
    assert client.get("/reports/licenses/expiring", params={"days": 60}).json() == []
    assert client.get("/reports/warranty/expiring").json() == []
    assert client.get("/reports/warranty/expiring", params={"days": -1}).status_code == 422


def test_report_as_csv_and_xlsx(client, factory):
    factory.equipment()
    response = client.get("/reports/inventory", params={"format": "csv"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert 'filename="inventory.csv"' in response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert rows[0]["count"] == "1"

    response = client.get("/reports/inventory", params={"format": "xlsx"})
    assert response.headers["content-type"].startswith("application/vnd.openxml")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[1]] == ["office", "type", "status", "count"]
    assert sheet.cell(2, 4).value == 1

    assert client.get("/reports/inventory", params={"format": "pdf"}).status_code == 422
