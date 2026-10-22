"""Reading and writing tables as CSV or XLSX."""

import csv
import io
from collections.abc import Iterable, Mapping, Sequence
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from typing import Any, Literal

from openpyxl import Workbook, load_workbook

from datakyt.services.errors import InvalidInputError

Format = Literal["csv", "xlsx"]

MEDIA_TYPES: dict[Format, str] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Cells starting with these characters are executed as formulas by spreadsheet apps.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _cell(value: Any) -> Any:
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    return value


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def write_table(rows: Iterable[Mapping[str, Any]], columns: Sequence[str], fmt: Format) -> bytes:
    """Serialize ``rows`` with the given column order."""
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([_text(_safe_text(_cell(row.get(c)))) for c in columns])
        return buffer.getvalue().encode("utf-8")

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("data")
    sheet.append(list(columns))
    for row in rows:
        sheet.append([_xlsx_cell(_safe_text(_cell(row.get(c)))) for c in columns])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _xlsx_cell(value: Any) -> Any:
    return float(value) if isinstance(value, Decimal) else value


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, date | datetime):
        return value.isoformat()
    return str(value)


def read_table(data: bytes, fmt: Format) -> list[dict[str, Any]]:
    """Parse an uploaded file into a list of rows keyed by the header row.

    Empty cells are omitted from the row so that model defaults apply.
    """
    if fmt == "csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InvalidInputError("CSV file must be UTF-8 encoded") from exc
        reader = csv.DictReader(io.StringIO(text))
        return [
            {k.strip(): v.strip() for k, v in row.items() if k and v is not None and v.strip()}
            for row in reader
        ]

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many unrelated types for bad files
        raise InvalidInputError("file is not a valid XLSX workbook") from exc
    sheet = workbook.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    names = [str(h).strip() if h is not None else "" for h in header]
    result: list[dict[str, Any]] = []
    for values in rows_iter:
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        row: dict[str, Any] = {}
        for name, value in zip(names, values, strict=False):
            if not name or value is None or (isinstance(value, str) and not value.strip()):
                continue
            if isinstance(value, datetime):
                value = value.date() if value.time() == datetime.min.time() else value
            row[name] = value.strip() if isinstance(value, str) else value
        result.append(row)
    workbook.close()
    return result
